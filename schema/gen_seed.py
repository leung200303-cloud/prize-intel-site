# -*- coding: utf-8 -*-
"""從 Excel 源頭清單 + module1/aliases.py 生成 02_seed.sql。

單一真相：
  - 源頭清單 → Excel（01_源頭清單_v0.3.xlsx）
  - 廠商／作品／別名 → module1/aliases.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
XLSX = HERE.parent / "01_源頭清單_v0.3.xlsx"
OUT = HERE / "02_seed.sql"

sys.path.insert(0, str(HERE.parent / "module1"))
from aliases import BRANDS, MAKERS, WORKS  # noqa: E402

CODE_BY_ID = {
    "A1": "bsp_prize", "A2": "ichiban_kuji", "A3": "sega_plaza", "A4": "taito_prize",
    "A5": "furyu", "A6": "eikoh", "A7": "system_service", "A8": "konami_ags",
    "A9": "kitan_club", "A10": "prtimes",
    "B1": "dengeki_prizefair", "B2": "dengeki_news", "B3": "am_journal",
    "B4": "hobby_watch", "B5": "nijimen", "B6": "medery",
    "C1": "kore_tore", "C2": "kaho_memo", "C3": "personal_blogs",
    "D1": "mercari", "D2": "yahoo_auction", "D3": "surugaya", "D4": "mandarake",
    "D5": "lashinbang", "D6": "paypay_flea", "D7": "amazon_jp",
    "E1": "b2b_channel", "E2": "hiroichi", "E3": "standby", "E4": "jaia",
    "E5": "prizefair_main",
}
MODE = {"自動定時": "auto_daily", "每週": "weekly", "事件驅動": "event_driven",
        "只放連結": "link_only", "手動補錄": "manual", "唔爬": "skip"}
PRI = {"P0": 0, "P1": 1, "P2": 2}

def q(v):
    return "NULL" if v is None else "'" + str(v).replace("'", "''") + "'"


def arr(values) -> str:
    uniq = list(dict.fromkeys(v for v in values if v))
    return "ARRAY[" + ", ".join(q(v) for v in uniq) + "]::text[]"


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["01_源頭清單"]
    src_rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        sid, layer, name, url, ctype, covers, freq, anti, mode, pri, ok, note = r[:12]
        code = CODE_BY_ID.get(sid)
        if not code:
            raise SystemExit(f"未登記 source code: {sid} {name}")
        if url in ("待確認", "非公開", None):
            url = None
        src_rows.append(
            f"({q(code)}, {q(layer)}, {q(name)}, {q(url)}, {q(ctype)}, {q(covers)}, "
            f"{q(freq)}, {q(anti)}, {q(MODE[mode])}, {PRI[pri]}, "
            f"{'true' if ok == '✅' else 'false'}, {q(note)})"
        )

    lines = [
        "-- =====================================================================",
        "-- 景品採購情報搜尋站 / Prize Intel — 02_seed.sql",
        "-- 自動生成，請勿手改：",
        "--   廠商／作品／別名 ← module1/aliases.py（單一真相）",
        "--   源頭清單        ← 01_源頭清單_v0.3.xlsx",
        "-- =====================================================================",
        "",
        "BEGIN;",
        "",
        f"-- 廠商 {len(MAKERS)} 間（aliases 解決「品牌名 vs 公司名」）",
        "INSERT INTO maker (code, name_jp, name_en, aliases, sort_order) VALUES",
        ",\n".join(
            f"  ({q(c)}, {q(jp)}, {q(en)}, {arr([*al, jp, en])}, {o})"
            for c, jp, en, al, o in MAKERS
        )
        + "\nON CONFLICT (code) DO NOTHING;",
        "",
        f"-- 系列線 {len(BRANDS)} 條（場主真正嘅搜尋入口）",
        "INSERT INTO brand_line (code, maker_code, name_jp, name_en, sort_order) VALUES",
        ",\n".join(f"  ({q(c)}, {q(m)}, {q(jp)}, {q(en)}, {o})"
                   for c, m, jp, en, o in BRANDS)
        + "\nON CONFLICT (code) DO NOTHING;",
        "",
        f"-- 作品 {len(WORKS)} 套（日文原名做 key，中英譯名／後綴變體入 aliases）",
        "INSERT INTO ip_work (code, title_jp, title_en, title_zh, aliases) VALUES",
        ",\n".join(
            f"  ({q(c)}, {q(jp)}, {q(en)}, {q(zh)}, {arr([*al, jp, en, zh])})"
            for c, jp, en, zh, al in WORKS
        )
        + "\nON CONFLICT (code) DO NOTHING;",
        "",
        f"-- 源頭 registry（{len(src_rows)} 條，對應 Excel 01_源頭清單）",
        "INSERT INTO source (code, layer, name_jp, url, content_type, covers,",
        "                    update_freq, anti_bot, crawl_mode, priority, verified, notes) VALUES",
        ",\n".join("  " + r for r in src_rows)
        + "\nON CONFLICT (code) DO NOTHING;",
        "",
        "COMMIT;",
        "",
    ]
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {OUT.name}: sources={len(src_rows)} makers={len(MAKERS)} "
          f"brands={len(BRANDS)} works={len(WORKS)}")


if __name__ == "__main__":
    main()
